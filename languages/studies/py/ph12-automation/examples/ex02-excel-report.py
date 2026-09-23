#!/usr/bin/env python3
# examples/ex02-excel-report.py —— Excel 自动化：openpyxl 生成设备状态周报
# 验证环境：Python 3.13.9，openpyxl 3.1.5（pip install openpyxl）
# 运行：python3 ex02-excel-report.py（离线可跑，已验证；xlsx 写入系统临时目录）
# 说明：对应主文档 3.2/3.5。建表头 → 写数据 → 合计行 → 样式（加粗/填充/冻结/列宽），
#       存盘后用 load_workbook 重新打开核对单元格——报表数字可审计。
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pathlib import Path
import tempfile

HEADERS = ["日期", "设备", "状态", "累计运行量(km)", "能耗(kWh)", "平均速度(km/h)"]

ROWS = [
    ("2026-09-01", "EV-001", "正常", 120.5, 18.2, 55.0),
    ("2026-09-01", "EV-002", "正常", 98.0, 15.1, 48.5),
    ("2026-09-01", "EV-003", "告警", 45.2, 7.8, 40.0),
    ("2026-09-02", "EV-001", "正常", 132.8, 19.6, 57.2),
    ("2026-09-02", "EV-002", "正常", 110.3, 16.4, 50.1),
    ("2026-09-02", "EV-003", "正常", 88.6, 13.9, 46.8),
    ("2026-09-03", "EV-001", "正常", 115.0, 17.5, 53.3),
    ("2026-09-03", "EV-002", "正常", 105.7, 15.9, 49.9),
    ("2026-09-03", "EV-003", "正常", 92.4, 14.2, 47.6),
]


def build_report(rows: list[tuple]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "设备状态周报"

    ws.append(HEADERS)  # 按行追加：表头
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    for cell in ws[1]:  # 表头样式：加粗白字 + 蓝底
        cell.font = header_font
        cell.fill = header_fill

    for r in rows:
        ws.append(r)  # 按行追加：数据

    # 合计行：累计运行量/能耗求和，平均速度求平均
    total_row = ws.max_row + 1
    ws.cell(total_row, 1, "合计")
    ws.cell(total_row, 4, round(sum(r[3] for r in rows), 1))
    ws.cell(total_row, 5, round(sum(r[4] for r in rows), 1))
    ws.cell(total_row, 6, round(sum(r[5] for r in rows) / len(rows), 2))
    for cell in ws[total_row]:
        cell.font = Font(bold=True)

    widths = [12, 10, 8, 12, 12, 16]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + col)].width = w  # 列宽
    ws.freeze_panes = "A2"  # 冻结表头行
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(horizontal="center")  # 日期列居中

    out = Path(tempfile.mkdtemp(prefix="ph12-ex02-")) / "device-weekly-report.xlsx"
    wb.save(out)
    return out


def main() -> None:
    out = build_report(ROWS)

    wb2 = load_workbook(out)  # 重新打开核对：数字可审计的关键一步
    ws2 = wb2["设备状态周报"]
    print("工作表:", ws2.title, "| 维度:", ws2.max_row, "行 ×", ws2.max_column, "列")
    print("表头:", " | ".join(str(c.value) for c in ws2[1]))
    total_row = ws2.max_row
    print("合计行:", ws2.cell(total_row, 1).value,
          "| 累计运行量:", ws2.cell(total_row, 4).value,
          "| 能耗:", ws2.cell(total_row, 5).value,
          "| 平均速度:", ws2.cell(total_row, 6).value)
    print("示例单元格 B2:", ws2["B2"].value, "| 冻结窗格:", ws2.freeze_panes)
    print("文件:", out, f"（{out.stat().st_size} 字节）")


if __name__ == "__main__":
    main()
